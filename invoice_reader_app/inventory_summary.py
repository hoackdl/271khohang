from django.shortcuts import render
from django.core.paginator import Paginator
from django.http import HttpResponse
import openpyxl
from invoice_reader_app.model_invoice import InvoiceItem, ProductName
from django.contrib.auth.decorators import login_required
from django.http import JsonResponse
from django.db.models import (
    Sum, F, Case, When, FloatField, Value, Q
)
from django.db.models.functions import Coalesce, Cast
from invoice_reader_app.models_purchaseorder import PurchaseOrder, PurchaseOrderItem  # đổi theo tên app của bạn
from invoice_reader_app.model_invoice import ProductName
from django.shortcuts import render, redirect
from django.contrib import messages
from django.db import transaction
from decimal import Decimal, InvalidOperation
from datetime import datetime
from .models_purchaseorder import PurchaseOrder, PurchaseOrderItem
from invoice_reader_app.model_invoice import ProductName, InventoryOpening


@login_required
def inventory_detail(request, sku):
    qs = (
        PurchaseOrderItem.objects
        .filter(sku_auto=sku)
        .select_related('purchase_order')
        .values(
            'purchase_order__po_number',
            'purchase_order__created_at',
            'purchase_order__phan_loai_phieu'
        )
        .annotate(
            so_luong=Sum(
                Coalesce(F('so_luong_quy_doi'), F('quantity'))
            ),
            thanh_tien=Sum('total_price')
        )
        .order_by('-purchase_order__created_at')
    )

    return JsonResponse(list(qs), safe=False)


def inventory_summary_export(request):
    qs = get_inventory_data_fast(request.GET.get('search'))

    wb = openpyxl.Workbook(write_only=True)
    ws = wb.create_sheet("XNT")

    ws.append([
        "SKU", "Tổng nhập", "Tổng xuất",
        "Tồn cuối", "Giá TB nhập", "Giá TB xuất"
    ])

    for row in qs:
        tong_nhap = row['tong_nhap']
        tong_xuat = row['tong_xuat']

        ws.append([
            row['sku'],
            tong_nhap,
            tong_xuat,
            tong_nhap - tong_xuat,
            row['gia_tb_nhap'],
            row['gia_tb_xuat'],
        ])

    response = HttpResponse(
        content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
    )
    response['Content-Disposition'] = 'attachment; filename=inventory.xlsx'
    wb.save(response)
    return response


from collections import defaultdict
from django.db.models import F

def get_inventory_data_fast(search=None):
    """
    Phiên bản tối ưu: xử lý nhanh hơn bằng cách:
    - Lấy ProductName một lần cho tất cả SKU.
    - Dùng dict để group SKU.
    """
    # 1. Lấy tất cả item nhập/xuất
    po_items_in = PurchaseOrderItem.objects.filter(
        purchase_order__phan_loai_phieu__in=['HH', 'PN']
    ).values(
        'sku', 'ten_goi_chung', 'unit', 'quantity', 'so_luong_quy_doi', 'total_price',
        'purchase_order__po_number', 'purchase_order__created_at'
    )

    po_items_out = PurchaseOrderItem.objects.filter(
        purchase_order__phan_loai_phieu='PX'
    ).values(
        'sku', 'ten_goi_chung', 'unit', 'quantity', 'total_price',
        'purchase_order__po_number', 'purchase_order__created_at'
    )

    # 2. Lấy tên gọi chung một lần cho tất cả SKU
    skus = set()
    for item in po_items_in:
        if item['sku']:
            skus.add(item['sku'])
    for item in po_items_out:
        if item['sku']:
            skus.add(item['sku'])

    product_map = dict(
        ProductName.objects.filter(sku__in=skus)
        .values_list('sku', 'ten_goi_chung')
    )

    # 3. Group dữ liệu
    grouped = defaultdict(lambda: {
        'sku': '',
        'ten_goi_chung': '',
        'dvt': '',
        'tong_nhap': 0,
        'tong_xuat': 0,
        'gia_tb_nhap': 0,
        'gia_tb_xuat': 0,
        'chi_tiet_nhap': [],
        'chi_tiet_xuat': [],
    })

    # --- nhập ---
    for item in po_items_in:
        sku = item['sku']
        if not sku:
            continue

        data = grouped[sku]
        data['sku'] = sku
        data['ten_goi_chung'] = product_map.get(sku, item['ten_goi_chung'])
        data['dvt'] = item['unit']

        qty = float(item['so_luong_quy_doi'] or item['quantity'] or 0)
        thanh_tien = float(item['total_price'] or 0)
        don_gia = thanh_tien / qty if qty else 0

        data['tong_nhap'] += qty
        data['chi_tiet_nhap'].append({
            'ma_po': item['purchase_order__po_number'],
            'so_luong': qty,
            'thanh_tien': float(item['total_price'] or 0),
            'don_gia': round(don_gia, 2),
            'ngay': item['purchase_order__created_at'],
        })

    # --- xuất ---
    for item in po_items_out:
        sku = item['sku']
        if not sku:
            continue

        data = grouped[sku]
        data['sku'] = sku
        data['ten_goi_chung'] = product_map.get(sku, item['ten_goi_chung'])
        data['dvt'] = item['unit']

        qty = float(item['quantity'] or 0)
        thanh_tien = float(item['total_price'] or 0)
        don_gia = thanh_tien / qty if qty else 0

        data['tong_xuat'] += qty
        data['chi_tiet_xuat'].append({
            'ma_po': item['purchase_order__po_number'],
            'so_luong': qty,
            'thanh_tien': float(item['total_price'] or 0),
            'don_gia': round(don_gia, 2),
            'ngay': item['purchase_order__created_at'],
        })

    # 4. Tính tồn cuối và giá TB
    data_list = []
    for sku, d in grouped.items():
        tong_tien_nhap = sum(
            c['thanh_tien'] for c in d['chi_tiet_nhap']
        )
        tong_tien_xuat = sum(
            c['thanh_tien'] for c in d['chi_tiet_xuat']
        )

        d['tong_tien_nhap'] = tong_tien_nhap
        d['tong_tien_xuat'] = tong_tien_xuat

        d['ton_cuoi'] = d['tong_nhap'] - d['tong_xuat']

        d['gia_tb_nhap'] = (
            tong_tien_nhap / d['tong_nhap']
            if d['tong_nhap'] else 0
        )

        d['gia_tb_xuat'] = (
            tong_tien_xuat / d['tong_xuat']
            if d['tong_xuat'] else 0
        )

        d['chi_tiet'] = (
            d['chi_tiet_nhap'] + d['chi_tiet_xuat']
        )

        data_list.append(d)

    # 4.5. Sắp xếp theo SKU
    data_list = sorted(data_list, key=lambda x: x['sku'] or '')
    
    # 5. Lọc theo search nếu cần
    if search:
        search_lower = search.lower()
        data_list = [
            d for d in data_list
            if search_lower in (d['sku'] or '').lower() or search_lower in (d['ten_goi_chung'] or '').lower()
        ]

    return data_list


from django.core.paginator import Paginator
from django.shortcuts import render
from django.contrib.auth.decorators import login_required


@login_required
def inventory_summary(request):
    search = request.GET.get('search', '').strip()
    per_page = int(request.GET.get('per_page', 20))
    page_number = request.GET.get('page')  # ✅ KHÔNG ÉP INT

    data_all = get_inventory_data_fast(search=search)

    # 🔥 đưa tồn âm lên đầu
    data_all.sort(key=lambda x: x['ton_cuoi'] >= 0)

    
    paginator = Paginator(data_all, per_page)
    page_obj = paginator.get_page(page_number)

    totals = {
        'tong_nhap': sum(d['tong_nhap'] for d in data_all),
        'tong_xuat': sum(d['tong_xuat'] for d in data_all),
        'ton_cuoi': sum(d['ton_cuoi'] for d in data_all),
        'gia_tb_nhap': round(
            sum(d['tong_nhap'] * d['gia_tb_nhap'] for d in data_all) /
            sum(d['tong_nhap'] for d in data_all)
        ) if sum(d['tong_nhap'] for d in data_all) else 0,
        'gia_tb_xuat': round(
            sum(d['tong_xuat'] * d['gia_tb_xuat'] for d in data_all) /
            sum(d['tong_xuat'] for d in data_all)
        ) if sum(d['tong_xuat'] for d in data_all) else 0,
    }

    return render(request, "inventory_summary.html", {
        "data": page_obj.object_list,
        "page_obj": page_obj,
        "search": search,
        "per_page": per_page,
        "totals": totals,
    })

from collections import defaultdict

def get_inventory_data_fast(
    search=None,
    date_from=None,
    date_to=None
):
    """
    Báo cáo Xuất - Nhập - Tồn

    Tồn cuối = Tồn đầu kỳ
               + Nhập trong khoảng ngày
               - Xuất trong khoảng ngày
    """

    # =========================================================
    # 1. TỒN ĐẦU KỲ
    # =========================================================
    opening_items = InventoryOpening.objects.all().values(
        'sku',
        'ten_goi_chung',
        'dvt',
        'so_luong',
        'don_gia',
        'thanh_tien',
        'ngay',
        'ghi_chu',
    )

    # =========================================================
    # 2. PHIẾU NHẬP
    # =========================================================
    po_items_in = PurchaseOrderItem.objects.filter(
        purchase_order__phan_loai_phieu__in=['HH', 'PN']
    )

    # Lọc ngày nhập
    if date_from:
        po_items_in = po_items_in.filter(
            purchase_order__created_at__date__gte=date_from
        )

    if date_to:
        po_items_in = po_items_in.filter(
            purchase_order__created_at__date__lte=date_to
        )

    po_items_in = po_items_in.values(
        'sku',
        'ten_goi_chung',
        'unit',
        'quantity',
        'so_luong_quy_doi',
        'total_price',
        'purchase_order__po_number',
        'purchase_order__created_at'
    )

    # =========================================================
    # 3. PHIẾU XUẤT
    # =========================================================
    po_items_out = PurchaseOrderItem.objects.filter(
        purchase_order__phan_loai_phieu='PX'
    )

    # Lọc ngày xuất
    if date_from:
        po_items_out = po_items_out.filter(
            purchase_order__created_at__date__gte=date_from
        )

    if date_to:
        po_items_out = po_items_out.filter(
            purchase_order__created_at__date__lte=date_to
        )

    po_items_out = po_items_out.values(
        'sku',
        'ten_goi_chung',
        'unit',
        'quantity',
        'total_price',
        'purchase_order__po_number',
        'purchase_order__created_at',

        'purchase_order__invoice__so_hoa_don',
        'purchase_order__invoice__ky_hieu',
    )

    # =========================================================
    # PHẦN CÒN LẠI GIỮ NGUYÊN
    # =========================================================


    # =========================================================
    # 4. LẤY TOÀN BỘ SKU
    # =========================================================
    skus = set()

    for item in opening_items:
        if item['sku']:
            skus.add(item['sku'])

    for item in po_items_in:
        if item['sku']:
            skus.add(item['sku'])

    for item in po_items_out:
        if item['sku']:
            skus.add(item['sku'])

    # =========================================================
    # 5. MAP TÊN SẢN PHẨM
    # =========================================================
    product_map = dict(
        ProductName.objects
        .filter(sku__in=skus)
        .values_list('sku', 'ten_goi_chung')
    )

    # =========================================================
    # 6. GROUP THEO SKU
    # =========================================================
    grouped = defaultdict(lambda: {
        'sku': '',
        'ten_goi_chung': '',
        'dvt': '',

        'ton_dau_ky': 0,
        'tong_nhap': 0,
        'tong_xuat': 0,

        'tong_tien_dau_ky': 0,
        'tong_tien_nhap': 0,
        'tong_tien_xuat': 0,

        'gia_tb_dau_ky': 0,
        'gia_tb_nhap': 0,
        'gia_tb_xuat': 0,

        'ton_cuoi': 0,

        'chi_tiet_dau_ky': [],
        'chi_tiet_nhap': [],
        'chi_tiet_xuat': [],
    })

    # =========================================================
    # 7. ĐẦU KỲ
    # =========================================================
    for item in opening_items:

        sku = item['sku']

        if not sku:
            continue

        data = grouped[sku]

        data['sku'] = sku

        data['ten_goi_chung'] = (
            product_map.get(sku)
            or item['ten_goi_chung']
            or ''
        )

        data['dvt'] = item['dvt'] or ''

        qty = float(item['so_luong'] or 0)
        don_gia = float(item['don_gia'] or 0)

        thanh_tien = float(
            item['thanh_tien']
            or (qty * don_gia)
            or 0
        )

        data['ton_dau_ky'] += qty
        data['tong_tien_dau_ky'] += thanh_tien

        data['chi_tiet_dau_ky'].append({
            'ma_po': 'ĐẦU KỲ',
            'so_luong': qty,
            'don_gia': don_gia,
            'thanh_tien': thanh_tien,
            'ngay': item['ngay'],
            'ghi_chu': item['ghi_chu'],
        })

    # =========================================================
    # 8. NHẬP
    # =========================================================
    for item in po_items_in:

        sku = item['sku']

        if not sku:
            continue

        data = grouped[sku]

        data['sku'] = sku

        data['ten_goi_chung'] = (
            product_map.get(sku)
            or item['ten_goi_chung']
            or ''
        )

        data['dvt'] = item['unit'] or ''

        qty = float(
            item['so_luong_quy_doi']
            or item['quantity']
            or 0
        )

        thanh_tien = float(
            item['total_price']
            or 0
        )

        don_gia = (
            thanh_tien / qty
            if qty else 0
        )

        data['tong_nhap'] += qty
        data['tong_tien_nhap'] += thanh_tien

        data['chi_tiet_nhap'].append({
            'ma_po': item['purchase_order__po_number'],
            'so_luong': qty,
            'thanh_tien': thanh_tien,
            'don_gia': round(don_gia, 2),
            'ngay': item['purchase_order__created_at'],
        })

    # =========================================================
    # 9. XUẤT
    # =========================================================
    # =========================================================
    # 9. XUẤT
    # =========================================================
    for item in po_items_out:

        sku = item['sku']

        if not sku:
            continue

        data = grouped[sku]

        data['sku'] = sku

        data['ten_goi_chung'] = (
            product_map.get(sku)
            or item['ten_goi_chung']
            or ''
        )

        data['dvt'] = item['unit'] or ''

        qty = float(item['quantity'] or 0)

        thanh_tien = float(item['total_price'] or 0)

        don_gia = (
            thanh_tien / qty
            if qty else 0
        )

        data['tong_xuat'] += qty
        data['tong_tien_xuat'] += thanh_tien

        # Kiểm tra PX đã tạo hóa đơn chưa
        so_hoa_don = item.get(
            'purchase_order__invoice__so_hoa_don'
        )

        ky_hieu = item.get(
            'purchase_order__invoice__ky_hieu'
        )

        da_tao_hoa_don = bool(so_hoa_don)

        data['chi_tiet_xuat'].append({
            'ma_po': item['purchase_order__po_number'],
            'so_luong': qty,
            'thanh_tien': thanh_tien,
            'don_gia': round(don_gia, 2),
            'ngay': item['purchase_order__created_at'],

            # Thông tin hóa đơn
            'da_tao_hoa_don': da_tao_hoa_don,
            'so_hoa_don': so_hoa_don or '',
            'ky_hieu': ky_hieu or '',
        })

    # =========================================================
    # SẮP XẾP CHI TIẾT PHIẾU XUẤT
    # =========================================================
    for sku, d in grouped.items():

        d['chi_tiet_xuat'].sort(
            key=lambda x: (
                # 1. PX chưa tạo hóa đơn lên trước
                0 if not x.get('da_tao_hoa_don') else 1,

                # 2. PX đã tạo → theo số hóa đơn tăng dần
                int(x.get('so_hoa_don') or 0)
                if str(x.get('so_hoa_don') or '').isdigit()
                else 0
            )
        )
    # =========================================================
    # 10. TÍNH TOÁN
    # =========================================================
    data_list = []

    for sku, d in grouped.items():

        d['gia_tb_dau_ky'] = (
            d['tong_tien_dau_ky'] / d['ton_dau_ky']
            if d['ton_dau_ky']
            else 0
        )

        d['gia_tb_nhap'] = (
            d['tong_tien_nhap'] / d['tong_nhap']
            if d['tong_nhap']
            else 0
        )

        d['gia_tb_xuat'] = (
            d['tong_tien_xuat'] / d['tong_xuat']
            if d['tong_xuat']
            else 0
        )

        # ⭐ CÔNG THỨC QUAN TRỌNG
        d['ton_cuoi'] = (
            d['ton_dau_ky']
            + d['tong_nhap']
            - d['tong_xuat']
        )

        d['chi_tiet'] = (
            d['chi_tiet_dau_ky']
            + d['chi_tiet_nhap']
            + d['chi_tiet_xuat']
        )

        data_list.append(d)

    # =========================================================
    # 11. SẮP XẾP
    # =========================================================
    data_list.sort(
        key=lambda x: x['sku'] or ''
    )

    # =========================================================
    # 12. SEARCH
    # =========================================================
    if search:

        search_lower = search.lower()

        data_list = [
            d for d in data_list
            if (
                search_lower in (d['sku'] or '').lower()
                or
                search_lower in (
                    d['ten_goi_chung'] or ''
                ).lower()
            )
        ]

    return data_list

@login_required
def inventory_summary(request):

    search = request.GET.get(
        'search',
        ''
    ).strip()

    try:
        per_page = int(
            request.GET.get(
                'per_page',
                20
            )
        )
    except (ValueError, TypeError):
        per_page = 20

    if per_page not in [10, 20, 50, 100]:
        per_page = 20

    page_number = request.GET.get('page')

    data_all = get_inventory_data_fast(
        search=search
    )

    # Đưa tồn âm lên đầu
    data_all.sort(
        key=lambda x: x['ton_cuoi'] >= 0
    )

    paginator = Paginator(
        data_all,
        per_page
    )

    page_obj = paginator.get_page(
        page_number
    )

    tong_nhap = sum(
        d['tong_nhap']
        for d in data_all
    )

    tong_xuat = sum(
        d['tong_xuat']
        for d in data_all
    )

    ton_dau_ky = sum(
        d['ton_dau_ky']
        for d in data_all
    )

    ton_cuoi = sum(
        d['ton_cuoi']
        for d in data_all
    )

    tong_tien_nhap = sum(
        d['tong_tien_nhap']
        for d in data_all
    )

    tong_tien_xuat = sum(
        d['tong_tien_xuat']
        for d in data_all
    )

    tong_tien_dau_ky = sum(
        d['tong_tien_dau_ky']
        for d in data_all
    )

    totals = {

        'ton_dau_ky': ton_dau_ky,

        'tong_nhap': tong_nhap,

        'tong_xuat': tong_xuat,

        'ton_cuoi': ton_cuoi,

        'gia_tb_dau_ky': round(
            tong_tien_dau_ky / ton_dau_ky
        ) if ton_dau_ky else 0,

        'gia_tb_nhap': round(
            tong_tien_nhap / tong_nhap
        ) if tong_nhap else 0,

        'gia_tb_xuat': round(
            tong_tien_xuat / tong_xuat
        ) if tong_xuat else 0,
    }

    return render(
        request,
        "inventory_summary.html",
        {
            "data": page_obj.object_list,
            "page_obj": page_obj,
            "search": search,
            "per_page": per_page,
            "totals": totals,
        }
    )


@login_required
def inventory_opening_import(request):

    # =========================================================
    # POST
    # =========================================================
    if request.method == 'POST':

        action = request.POST.get('action')

        # =====================================================
        # 1. LƯU THỦ CÔNG
        # =====================================================
        if action == 'save':

            skus = request.POST.getlist('sku[]')
            names = request.POST.getlist('ten_goi_chung[]')
            dvts = request.POST.getlist('dvt[]')
            quantities = request.POST.getlist('so_luong[]')
            prices = request.POST.getlist('don_gia[]')
            dates = request.POST.getlist('ngay[]')
            notes = request.POST.getlist('ghi_chu[]')

            count = 0

            with transaction.atomic():

                for i, sku in enumerate(skus):

                    sku = sku.strip()

                    if not sku:
                        continue

                    try:
                        qty = float(
                            quantities[i] or 0
                        )
                    except:
                        qty = 0

                    try:
                        price = float(
                            prices[i] or 0
                        )
                    except:
                        price = 0

                    if qty == 0:
                        continue

                    ngay = None

                    if i < len(dates) and dates[i]:

                        try:
                            ngay = datetime.strptime(
                                dates[i],
                                '%Y-%m-%d'
                            ).date()
                        except:
                            ngay = None

                    InventoryOpening.objects.create(
                        sku=sku,
                        ten_goi_chung=(
                            names[i]
                            if i < len(names)
                            else ''
                        ),
                        dvt=(
                            dvts[i]
                            if i < len(dvts)
                            else ''
                        ),
                        so_luong=qty,
                        don_gia=price,
                        thanh_tien=qty * price,
                        ngay=ngay,
                        ghi_chu=(
                            notes[i]
                            if i < len(notes)
                            else ''
                        ),
                    )

                    count += 1

            messages.success(
                request,
                f'Đã lưu {count} mặt hàng đầu kỳ.'
            )

            return redirect(
                'inventory_summary'
            )

        # =====================================================
        # 2. IMPORT EXCEL
        # =====================================================
        if action == 'excel':

            excel_file = request.FILES.get(
                'excel_file'
            )

            if not excel_file:

                messages.error(
                    request,
                    'Bạn chưa chọn file Excel.'
                )

                return redirect(
                    'inventory_opening_import'
                )

            try:

                wb = openpyxl.load_workbook(
                    excel_file,
                    data_only=True
                )

                ws = wb.active

                rows = list(
                    ws.iter_rows(
                        values_only=True
                    )
                )

                if not rows:

                    messages.error(
                        request,
                        'File Excel không có dữ liệu.'
                    )

                    return redirect(
                        'inventory_opening_import'
                    )

                # ---------------------------------------------
                # Xác định cột theo tên header
                # ---------------------------------------------
                headers = [
                    str(x).strip().lower()
                    if x is not None
                    else ''
                    for x in rows[0]
                ]

                def find_col(*names):

                    for name in names:

                        name = name.lower()

                        for i, header in enumerate(headers):

                            if header == name:
                                return i

                    return None

                col_sku = find_col(
                    'sku',
                    'mã sku',
                    'mã hàng'
                )

                col_name = find_col(
                    'ten_goi_chung',
                    'tên gọi chung',
                    'tên hàng'
                )

                # ĐVT: hỗ trợ cả tên cột cũ "unit"
                # và tên mới "dvt", "đvt"
                col_dvt = find_col(
                    'dvt',
                    'đvt',
                    'unit'
                )

                col_qty = find_col(
                    'so_luong',
                    'số lượng',
                    'sl',
                    'tồn đầu'
                )

                col_price = find_col(
                    'don_gia',
                    'đơn giá',
                    'giá'
                )

                col_date = find_col(
                    'ngay',
                    'ngày'
                )

                col_note = find_col(
                    'ghi_chu',
                    'ghi chú'
                )

                if col_sku is None:

                    messages.error(
                        request,
                        'Excel phải có cột SKU.'
                    )

                    return redirect(
                        'inventory_opening_import'
                    )

                if col_qty is None:

                    messages.error(
                        request,
                        'Excel phải có cột Số lượng.'
                    )

                    return redirect(
                        'inventory_opening_import'
                    )

                count = 0

                with transaction.atomic():

                    for row in rows[1:]:

                        if not row:
                            continue

                        sku = (
                            str(row[col_sku]).strip()
                            if col_sku < len(row)
                            and row[col_sku] is not None
                            else ''
                        )

                        if not sku:
                            continue

                        # -------------------------------------
                        # Tên
                        # -------------------------------------
                        name = ''

                        if (
                            col_name is not None
                            and col_name < len(row)
                            and row[col_name] is not None
                        ):
                            name = str(
                                row[col_name]
                            ).strip()

                        # -------------------------------------
                        # ĐVT
                        # -------------------------------------
                        dvt = ''

                        if (
                            col_dvt is not None
                            and col_dvt < len(row)
                            and row[col_dvt] is not None
                        ):
                            dvt = str(
                                row[col_dvt]
                            ).strip()

                        # -------------------------------------
                        # Số lượng
                        # -------------------------------------
                        try:

                            qty = float(
                                row[col_qty]
                                if col_qty < len(row)
                                and row[col_qty] is not None
                                else 0
                            )

                        except:
                            qty = 0

                        if qty == 0:
                            continue

                        # -------------------------------------
                        # Đơn giá
                        # -------------------------------------
                        try:

                            price = float(
                                row[col_price]
                                if col_price is not None
                                and col_price < len(row)
                                and row[col_price] is not None
                                else 0
                            )

                        except:
                            price = 0

                        # -------------------------------------
                        # Ngày
                        # -------------------------------------
                        ngay = None

                        if (
                            col_date is not None
                            and col_date < len(row)
                            and row[col_date]
                        ):

                            value = row[col_date]

                            if isinstance(
                                value,
                                datetime
                            ):
                                ngay = value.date()

                            elif hasattr(
                                value,
                                'date'
                            ):
                                ngay = value.date()

                            else:

                                try:
                                    ngay = datetime.strptime(
                                        str(value),
                                        '%d/%m/%Y'
                                    ).date()

                                except:

                                    try:
                                        ngay = datetime.strptime(
                                            str(value),
                                            '%Y-%m-%d'
                                        ).date()

                                    except:
                                        ngay = None

                        # -------------------------------------
                        # Ghi chú
                        # -------------------------------------
                        note = ''

                        if (
                            col_note is not None
                            and col_note < len(row)
                            and row[col_note] is not None
                        ):
                            note = str(
                                row[col_note]
                            )

                        InventoryOpening.objects.create(
                            sku=sku,
                            ten_goi_chung=name,
                            dvt=dvt,
                            so_luong=qty,
                            don_gia=price,
                            thanh_tien=qty * price,
                            ngay=ngay,
                            ghi_chu=note,
                        )

                        count += 1

                messages.success(
                    request,
                    f'Đã import {count} mặt hàng đầu kỳ.'
                )

            except Exception as e:

                messages.error(
                    request,
                    f'Lỗi đọc Excel: {e}'
                )

            return redirect(
                'inventory_summary'
            )

    # =========================================================
    # GET
    # =========================================================
    return render(
        request,
        'inventory_opening_import.html'
    )

from django.core.paginator import Paginator
from django.db.models import Q
from django.contrib.auth.decorators import login_required
from django.shortcuts import render

from invoice_reader_app.model_invoice import InventoryOpening


@login_required
def inventory_opening_list(request):

    search = request.GET.get("search", "").strip()

    per_page = request.GET.get("per_page", "20")

    try:
        per_page = int(per_page)
    except ValueError:
        per_page = 20

    qs = InventoryOpening.objects.all().order_by(
        "-ngay",
        "-id"
    )

    if search:

        qs = qs.filter(
            Q(sku__icontains=search) |
            Q(ten_goi_chung__icontains=search) |
            Q(ghi_chu__icontains=search)
        )

    paginator = Paginator(
        qs,
        per_page
    )

    page_number = request.GET.get("page")

    page_obj = paginator.get_page(
        page_number
    )

    return render(
        request,
        "inventory_opening_list.html",
        {
            "page_obj": page_obj,
            "search": search,
            "per_page": per_page,
        }
    )

from django.shortcuts import render, redirect, get_object_or_404

@login_required
@transaction.atomic
def inventory_opening_edit(request, pk):

    opening = get_object_or_404(
        InventoryOpening,
        pk=pk
    )

    if request.method == "POST":

        opening.sku = request.POST.get("sku", "").strip()
        opening.ten_goi_chung = request.POST.get(
            "ten_goi_chung", ""
        ).strip()

        opening.unit = request.POST.get(
            "dvt", ""
        ).strip()

        opening.so_luong = Decimal(
            request.POST.get("so_luong") or 0
        )

        opening.don_gia = Decimal(
            request.POST.get("don_gia") or 0
        )

        opening.thanh_tien = (
            opening.so_luong *
            opening.don_gia
        )

        ngay = request.POST.get("ngay")

        if ngay:
            opening.ngay = ngay
        else:
            opening.ngay = None

        opening.ghi_chu = request.POST.get(
            "ghi_chu", ""
        ).strip()

        opening.save()

        messages.success(
            request,
            f"Đã cập nhật phiếu đầu kỳ SKU {opening.sku}"
        )

        return redirect("inventory_opening_list")

    return render(
        request,
        "inventory_opening_edit.html",
        {
            "opening": opening,
        }
    )


@login_required
@transaction.atomic
def inventory_opening_delete(request, pk):

    opening = get_object_or_404(
        InventoryOpening,
        pk=pk
    )

    if request.method == "POST":

        sku = opening.sku

        opening.delete()

        messages.success(
            request,
            f"Đã xóa hàng đầu kỳ {sku}"
        )

    return redirect("inventory_opening_list")


from datetime import datetime, date
from collections import defaultdict
from datetime import datetime

def get_inventory_data_fast(
    search=None,
    date_from=None,
    date_to=None
):
    """
    BÁO CÁO XUẤT - NHẬP - TỒN THEO KHOẢNG THỜI GIAN

    Tồn đầu khoảng thời gian
        = Tồn đầu năm
        + Nhập trước ngày bắt đầu
        - Xuất trước ngày bắt đầu

    Nhập kỳ
        = Nhập từ ngày bắt đầu đến ngày kết thúc

    Xuất kỳ
        = Xuất từ ngày bắt đầu đến ngày kết thúc

    Tồn cuối
        = Tồn đầu khoảng thời gian
        + Nhập kỳ
        - Xuất kỳ
    """
    # =========================================================
    # CHUYỂN STRING -> DATE
    # =========================================================

    if isinstance(date_from, str) and date_from:
        try:
            date_from = datetime.strptime(
                date_from,
                "%Y-%m-%d"
            ).date()
        except ValueError:
            date_from = None

    if isinstance(date_to, str) and date_to:
        try:
            date_to = datetime.strptime(
                date_to,
                "%Y-%m-%d"
            ).date()
        except ValueError:
            date_to = None

            
    # =========================================================
    # 1. TỒN ĐẦU NĂM
    # =========================================================

    opening_items = (
        InventoryOpening.objects
        .all()
        .values(
            'sku',
            'ten_goi_chung',
            'dvt',
            'so_luong',
            'don_gia',
            'thanh_tien',
            'ngay',
            'ghi_chu',
        )
    )

    # =========================================================
    # 2. PHIẾU NHẬP
    # =========================================================

    po_items_in = (
        PurchaseOrderItem.objects
        .filter(
            purchase_order__phan_loai_phieu__in=[
                'HH',
                'PN'
            ]
        )
        .values(
            'sku',
            'ten_goi_chung',
            'unit',
            'quantity',
            'so_luong_quy_doi',
            'total_price',
            'purchase_order__po_number',
            'purchase_order__created_at'
        )
    )

    # =========================================================
    # 3. PHIẾU XUẤT
    # =========================================================

    po_items_out = (
        PurchaseOrderItem.objects
        .filter(
            purchase_order__phan_loai_phieu='PX'
        )
        .values(
            'sku',
            'ten_goi_chung',
            'unit',
            'quantity',
            'total_price',
            'purchase_order__po_number',
            'purchase_order__created_at',

            # Thông tin hóa đơn
            'purchase_order__invoice__so_hoa_don',
            'purchase_order__invoice__ky_hieu',
        )
    )

    # =========================================================
    # 4. LẤY TOÀN BỘ SKU
    # =========================================================

    skus = set()

    for item in opening_items:
        if item['sku']:
            skus.add(item['sku'])

    for item in po_items_in:
        if item['sku']:
            skus.add(item['sku'])

    for item in po_items_out:
        if item['sku']:
            skus.add(item['sku'])

    # =========================================================
    # 5. MAP TÊN SẢN PHẨM
    # =========================================================

    product_map = dict(
        ProductName.objects
        .filter(sku__in=skus)
        .values_list(
            'sku',
            'ten_goi_chung'
        )
    )

    # =========================================================
    # 6. GROUP THEO SKU
    # =========================================================

    grouped = defaultdict(lambda: {

        'sku': '',
        'ten_goi_chung': '',
        'dvt': '',

        # -----------------------------------------------------
        # TỒN ĐẦU NĂM
        # -----------------------------------------------------

        'ton_dau_nam': 0,
        'tong_tien_dau_nam': 0,

        # -----------------------------------------------------
        # NHẬP TRƯỚC KỲ
        # -----------------------------------------------------

        'nhap_truoc_ky': 0,
        'tien_nhap_truoc_ky': 0,

        # -----------------------------------------------------
        # XUẤT TRƯỚC KỲ
        # -----------------------------------------------------

        'xuat_truoc_ky': 0,
        'tien_xuat_truoc_ky': 0,

        # -----------------------------------------------------
        # TỒN ĐẦU KHOẢNG THỜI GIAN
        # -----------------------------------------------------

        'ton_dau_ky': 0,
        'tong_tien_dau_ky': 0,

        # -----------------------------------------------------
        # NHẬP TRONG KỲ
        # -----------------------------------------------------

        'tong_nhap': 0,
        'tong_tien_nhap': 0,

        # -----------------------------------------------------
        # XUẤT TRONG KỲ
        # -----------------------------------------------------

        'tong_xuat': 0,
        'tong_tien_xuat': 0,

        # -----------------------------------------------------
        # TỒN CUỐI
        # -----------------------------------------------------

        'ton_cuoi': 0,

        # -----------------------------------------------------
        # GIÁ
        # -----------------------------------------------------

        'gia_tb_dau_ky': 0,
        'gia_tb_nhap': 0,
        'gia_tb_xuat': 0,

        # -----------------------------------------------------
        # CHI TIẾT
        # -----------------------------------------------------

        'chi_tiet_dau_ky': [],
        'chi_tiet_nhap': [],
        'chi_tiet_xuat': [],
    })

    # =========================================================
    # 7. TỒN ĐẦU NĂM
    # =========================================================

    for item in opening_items:

        sku = item['sku']

        if not sku:
            continue

        data = grouped[sku]

        data['sku'] = sku

        data['ten_goi_chung'] = (
            product_map.get(sku)
            or item['ten_goi_chung']
            or ''
        )

        data['dvt'] = item['dvt'] or ''

        qty = float(
            item['so_luong'] or 0
        )

        don_gia = float(
            item['don_gia'] or 0
        )

        thanh_tien = float(
            item['thanh_tien']
            or (qty * don_gia)
            or 0
        )

        data['ton_dau_nam'] += qty

        data['tong_tien_dau_nam'] += thanh_tien

        # Nếu không lọc ngày thì vẫn giữ chi tiết đầu năm
        data['chi_tiet_dau_ky'].append({
            'ma_po': 'ĐẦU NĂM',
            'so_luong': qty,
            'don_gia': don_gia,
            'thanh_tien': thanh_tien,
            'ngay': item['ngay'],
            'ghi_chu': item['ghi_chu'],
        })

    # =========================================================
    # 8. XỬ LÝ PHIẾU NHẬP
    # =========================================================

    for item in po_items_in:

        sku = item['sku']

        if not sku:
            continue

        data = grouped[sku]

        data['sku'] = sku

        data['ten_goi_chung'] = (
            product_map.get(sku)
            or item['ten_goi_chung']
            or ''
        )

        data['dvt'] = item['unit'] or ''

        qty = float(
            item['so_luong_quy_doi']
            or item['quantity']
            or 0
        )

        thanh_tien = float(
            item['total_price']
            or 0
        )

        don_gia = (
            thanh_tien / qty
            if qty
            else 0
        )

        created_at = item[
            'purchase_order__created_at'
        ]

        transaction_date = (
            created_at.date()
            if created_at
            and hasattr(created_at, 'date')
            else created_at
        )

        # =====================================================
        # NHẬP TRƯỚC NGÀY BẮT ĐẦU
        # =====================================================

        is_before_period = (
            date_from
            and transaction_date
            and transaction_date < date_from
        )

        # =====================================================
        # NHẬP TRONG KỲ
        # =====================================================

        is_in_period = True

        if date_from:
            is_in_period = (
                transaction_date
                and transaction_date >= date_from
            )

        if date_to:
            is_in_period = (
                is_in_period
                and transaction_date
                and transaction_date <= date_to
            )

        # -----------------------------------------------------
        # NHẬP TRƯỚC KỲ
        # -----------------------------------------------------

        if is_before_period:

            data['nhap_truoc_ky'] += qty

            data['tien_nhap_truoc_ky'] += thanh_tien

        # -----------------------------------------------------
        # NHẬP TRONG KỲ
        # -----------------------------------------------------

        elif is_in_period:

            data['tong_nhap'] += qty

            data['tong_tien_nhap'] += thanh_tien

            data['chi_tiet_nhap'].append({
                'ma_po': item[
                    'purchase_order__po_number'
                ],
                'so_luong': qty,
                'thanh_tien': thanh_tien,
                'don_gia': round(
                    don_gia,
                    2
                ),
                'ngay': created_at,
            })

    # =========================================================
    # 9. XỬ LÝ PHIẾU XUẤT
    # =========================================================

    for item in po_items_out:

        sku = item['sku']

        if not sku:
            continue

        data = grouped[sku]

        data['sku'] = sku

        data['ten_goi_chung'] = (
            product_map.get(sku)
            or item['ten_goi_chung']
            or ''
        )

        data['dvt'] = item['unit'] or ''

        qty = float(
            item['quantity'] or 0
        )

        thanh_tien = float(
            item['total_price'] or 0
        )

        don_gia = (
            thanh_tien / qty
            if qty
            else 0
        )

        created_at = item[
            'purchase_order__created_at'
        ]

        transaction_date = (
            created_at.date()
            if created_at
            and hasattr(created_at, 'date')
            else created_at
        )

        # =====================================================
        # XUẤT TRƯỚC KỲ
        # =====================================================

        is_before_period = (
            date_from
            and transaction_date
            and transaction_date < date_from
        )

        # =====================================================
        # XUẤT TRONG KỲ
        # =====================================================

        is_in_period = True

        if date_from:
            is_in_period = (
                transaction_date
                and transaction_date >= date_from
            )

        if date_to:
            is_in_period = (
                is_in_period
                and transaction_date
                and transaction_date <= date_to
            )

        # -----------------------------------------------------
        # XUẤT TRƯỚC KỲ
        # -----------------------------------------------------

        if is_before_period:

            data['xuat_truoc_ky'] += qty

            data['tien_xuat_truoc_ky'] += thanh_tien

        # -----------------------------------------------------
        # XUẤT TRONG KỲ
        # -----------------------------------------------------

        elif is_in_period:

            data['tong_xuat'] += qty

            data['tong_tien_xuat'] += thanh_tien

            so_hoa_don = item.get(
                'purchase_order__invoice__so_hoa_don'
            )

            ky_hieu = item.get(
                'purchase_order__invoice__ky_hieu'
            )

            data['chi_tiet_xuat'].append({

                'ma_po': item[
                    'purchase_order__po_number'
                ],

                'so_luong': qty,

                'thanh_tien': thanh_tien,

                'don_gia': round(
                    don_gia,
                    2
                ),

                'ngay': created_at,

                'da_tao_hoa_don': bool(
                    so_hoa_don
                ),

                'so_hoa_don': (
                    so_hoa_don or ''
                ),

                'ky_hieu': (
                    ky_hieu or ''
                ),
            })

    # =========================================================
    # 10. TÍNH TỒN ĐẦU KHOẢNG THỜI GIAN
    # =========================================================

    for sku, data in grouped.items():

        data['ton_dau_ky'] = (
            data['ton_dau_nam']
            + data['nhap_truoc_ky']
            - data['xuat_truoc_ky']
        )

        data['tong_tien_dau_ky'] = (
            data['tong_tien_dau_nam']
            + data['tien_nhap_truoc_ky']
            - data['tien_xuat_truoc_ky']
        )

        # =====================================================
        # GIÁ TB TỒN ĐẦU KỲ
        # =====================================================

        data['gia_tb_dau_ky'] = (
            data['tong_tien_dau_ky']
            / data['ton_dau_ky']
            if data['ton_dau_ky']
            else 0
        )

        # =====================================================
        # GIÁ TB NHẬP TRONG KỲ
        # =====================================================

        data['gia_tb_nhap'] = (
            data['tong_tien_nhap']
            / data['tong_nhap']
            if data['tong_nhap']
            else 0
        )

        # =====================================================
        # GIÁ TB XUẤT TRONG KỲ
        # =====================================================

        data['gia_tb_xuat'] = (
            data['tong_tien_xuat']
            / data['tong_xuat']
            if data['tong_xuat']
            else 0
        )

        # =====================================================
        # TỒN CUỐI
        # =====================================================

        data['ton_cuoi'] = (
            data['ton_dau_ky']
            + data['tong_nhap']
            - data['tong_xuat']
        )

        # =====================================================
        # CHI TIẾT
        # =====================================================

        data['chi_tiet'] = (
            data['chi_tiet_dau_ky']
            + data['chi_tiet_nhap']
            + data['chi_tiet_xuat']
        )

    # =========================================================
    # 11. SẮP XẾP PX
    # =========================================================

    for sku, data in grouped.items():

        data['chi_tiet_xuat'].sort(
            key=lambda x: (
                0
                if not x.get('da_tao_hoa_don')
                else 1,

                int(
                    x.get('so_hoa_don') or 0
                )
                if str(
                    x.get('so_hoa_don') or ''
                ).isdigit()
                else 0
            )
        )

    # =========================================================
    # 12. DATA LIST
    # =========================================================

    data_list = list(
        grouped.values()
    )

    # =========================================================
    # 13. SEARCH
    # =========================================================

    if search:

        search_lower = search.lower()

        data_list = [
            d
            for d in data_list
            if (
                search_lower
                in (d['sku'] or '').lower()
                or
                search_lower
                in (
                    d['ten_goi_chung'] or ''
                ).lower()
            )
        ]

    # =========================================================
    # 14. SẮP XẾP SKU
    # =========================================================

    data_list.sort(
        key=lambda x: x['sku'] or ''
    )

    return data_list



@login_required
def inventory_summary(request):

    search = request.GET.get(
        'search',
        ''
    ).strip()

    date_from = request.GET.get(
        'date_from',
        ''
    ).strip()

    date_to = request.GET.get(
        'date_to',
        ''
    ).strip()

    # =========================================================
    # PARSE NGÀY
    # =========================================================

    parsed_date_from = None
    parsed_date_to = None

    if date_from:

        try:
            parsed_date_from = datetime.strptime(
                date_from,
                '%Y-%m-%d'
            ).date()

        except ValueError:
            date_from = ''

    if date_to:

        try:
            parsed_date_to = datetime.strptime(
                date_to,
                '%Y-%m-%d'
            ).date()

        except ValueError:
            date_to = ''

    # Nếu nhập ngược ngày thì tự đổi
    if (
        parsed_date_from
        and parsed_date_to
        and parsed_date_from > parsed_date_to
    ):

        parsed_date_from, parsed_date_to = (
            parsed_date_to,
            parsed_date_from
        )

        date_from = parsed_date_from.strftime(
            '%Y-%m-%d'
        )

        date_to = parsed_date_to.strftime(
            '%Y-%m-%d'
        )

    # =========================================================
    # PHÂN TRANG
    # =========================================================

    try:

        per_page = int(
            request.GET.get(
                'per_page',
                20
            )
        )

    except (ValueError, TypeError):

        per_page = 20

    if per_page not in [
        10,
        20,
        50,
        100
    ]:

        per_page = 20

    page_number = request.GET.get(
        'page'
    )

    # =========================================================
    # LẤY DỮ LIỆU XNT
    # =========================================================

    data_all = get_inventory_data_fast(
        search=search,
        date_from=parsed_date_from,
        date_to=parsed_date_to
    )

    # Tồn âm đưa lên đầu
    data_all.sort(
        key=lambda x: x['ton_cuoi'] >= 0
    )

    # =========================================================
    # PAGINATION
    # =========================================================

    paginator = Paginator(
        data_all,
        per_page
    )

    page_obj = paginator.get_page(
        page_number
    )

    # =========================================================
    # TỔNG
    # =========================================================

    tong_nhap = sum(
        d['tong_nhap']
        for d in data_all
    )

    tong_xuat = sum(
        d['tong_xuat']
        for d in data_all
    )

    ton_dau_ky = sum(
        d['ton_dau_ky']
        for d in data_all
    )

    ton_cuoi = sum(
        d['ton_cuoi']
        for d in data_all
    )

    tong_tien_nhap = sum(
        d['tong_tien_nhap']
        for d in data_all
    )

    tong_tien_xuat = sum(
        d['tong_tien_xuat']
        for d in data_all
    )

    tong_tien_dau_ky = sum(
        d['tong_tien_dau_ky']
        for d in data_all
    )

    # =========================================================
    # TOTALS
    # =========================================================

    totals = {

        'ton_dau_ky': ton_dau_ky,

        'tong_nhap': tong_nhap,

        'tong_xuat': tong_xuat,

        'ton_cuoi': ton_cuoi,

        'gia_tb_dau_ky': round(
            tong_tien_dau_ky
            / ton_dau_ky
        ) if ton_dau_ky else 0,

        'gia_tb_nhap': round(
            tong_tien_nhap
            / tong_nhap
        ) if tong_nhap else 0,

        'gia_tb_xuat': round(
            tong_tien_xuat
            / tong_xuat
        ) if tong_xuat else 0,
    }
    query_params = request.GET.copy()
    query_params.pop('page', None)

    query_params = query_params.urlencode()

    return render(
        request,
        "inventory_summary.html",
        {
            "data": page_obj.object_list,
            "page_obj": page_obj,

            "search": search,
            "date_from": date_from,
            "date_to": date_to,
            "per_page": per_page,

            "query_params": query_params,

            "totals": totals,
        }
    )




@login_required
def inventory_summary_export(request):

    # =========================================================
    # BỘ LỌC
    # =========================================================
    search = request.GET.get(
        'search',
        ''
    ).strip()

    date_from = request.GET.get(
        'date_from',
        ''
    ).strip()

    date_to = request.GET.get(
        'date_to',
        ''
    ).strip()

    # =========================================================
    # LẤY DỮ LIỆU
    # =========================================================
    qs = get_inventory_data_fast(
        search=search,
        date_from=date_from,
        date_to=date_to,
    )

    # =========================================================
    # TẠO FILE EXCEL
    # =========================================================
    wb = openpyxl.Workbook(
        write_only=True
    )

    ws = wb.create_sheet(
        "XNT"
    )

    # =========================================================
    # TIÊU ĐỀ
    # =========================================================

    # Tiêu đề chính
    ws.append([
        "BẢNG NHẬP - XUẤT - TỒN"
    ])

    # Khoảng thời gian
    if date_from and date_to:
        ws.append([
            f"Từ ngày {date_from} đến ngày {date_to}"
        ])

    elif date_from:
        ws.append([
            f"Từ ngày {date_from}"
        ])

    elif date_to:
        ws.append([
            f"Đến ngày {date_to}"
        ])

    else:
        ws.append([
            "Toàn bộ thời gian"
        ])

    # Dòng trống
    ws.append([])

    # =========================================================
    # HEADER
    # =========================================================
    ws.append([
        "SKU",
        "Tên gọi chung",
        "ĐVT",
        "Tồn đầu kỳ",
        "Tổng nhập",
        "Tổng xuất",
        "Tồn cuối",
        "Giá TB nhập",
        "Giá TB xuất",
    ])

    # =========================================================
    # DATA
    # =========================================================
    for row in qs:

        tong_nhap = row['tong_nhap']
        tong_xuat = row['tong_xuat']

        ton_dau_ky = row['ton_dau_ky']

        ton_cuoi = (
            ton_dau_ky
            + tong_nhap
            - tong_xuat
        )

        ws.append([
            row['sku'],
            row.get(
                'ten_goi_chung',
                ''
            ),
            row.get(
                'dvt',
                ''
            ),
            ton_dau_ky,
            tong_nhap,
            tong_xuat,
            ton_cuoi,
            row['gia_tb_nhap'],
            row['gia_tb_xuat'],
        ])

    # =========================================================
    # RESPONSE
    # =========================================================
    response = HttpResponse(
        content_type=(
            'application/vnd.openxmlformats-officedocument.'
            'spreadsheetml.sheet'
        )
    )

    # Tên file
    if date_from and date_to:
        filename = (
            f"nhap_xuat_ton_"
            f"{date_from}_den_{date_to}.xlsx"
        )

    elif date_from:
        filename = (
            f"nhap_xuat_ton_tu_{date_from}.xlsx"
        )

    elif date_to:
        filename = (
            f"nhap_xuat_ton_den_{date_to}.xlsx"
        )

    else:
        filename = "nhap_xuat_ton.xlsx"

    response['Content-Disposition'] = (
        f'attachment; filename="{filename}"'
    )

    wb.save(response)

    return response