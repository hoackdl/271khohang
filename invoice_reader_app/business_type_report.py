from datetime import datetime

from django.shortcuts import render
from django.core.paginator import Paginator
from invoice_reader_app.model_invoice import Invoice
from invoice_reader_app.models_activitytype import ActivityType
from invoice_reader_app.models_purchaseorder import PurchaseOrder

def business_type_report(request):


    # =====================================================
    # NĂM TÀI CHÍNH
    # =====================================================

    current_year = request.session.get(
        "fiscal_year",
        datetime.now().year
    )

    year_param = request.GET.get("year", "").strip()

    # Nếu chọn All
    if year_param.lower() == "all":
        selected_year = "all"

    # Nếu chọn năm cụ thể
    elif year_param:
        try:
            selected_year = int(year_param)
            request.session["fiscal_year"] = selected_year
            current_year = selected_year
        except (ValueError, TypeError):
            selected_year = current_year

    else:
        selected_year = current_year

    # =====================================================
    # LOẠI HÌNH KINH DOANH
    # =====================================================
    activity_types = (
        ActivityType.objects
        .filter(is_active=True)
        .order_by("code")
    )

    selected_activity_type = request.GET.get(
        "activity_type",
        ""
    ).strip()

    # =====================================================
    # NGÀY LỌC
    # =====================================================
    start_date = request.GET.get(
        "start_date",
        ""
    ).strip()

    end_date = request.GET.get(
        "end_date",
        ""
    ).strip()

    # =====================================================
    # QUERY CHUNG
    # =====================================================

    base_queryset = (
        PurchaseOrder.objects
        .select_related(
            "invoice",
            "activity_type"
        )
        .filter(
            invoice__ngay_hd__isnull=False,
            phan_loai_phieu__in=[
                "HH",
                "PN",
                "PX",
            ],
        )
    )

    # Chỉ lọc năm khi không chọn All
    if selected_year != "all":
        base_queryset = base_queryset.filter(
            invoice__fiscal_year=selected_year
        )

    # =====================================================
    # LỌC LOẠI HÌNH KINH DOANH
    # =====================================================
    if selected_activity_type:
        base_queryset = base_queryset.filter(
            activity_type_id=selected_activity_type
        )

    # =====================================================
    # LỌC TỪ NGÀY
    # =====================================================
    if start_date:
        base_queryset = base_queryset.filter(
            invoice__ngay_hd__gte=start_date
        )

    # =====================================================
    # LỌC ĐẾN NGÀY
    # =====================================================
    if end_date:
        base_queryset = base_queryset.filter(
            invoice__ngay_hd__lte=end_date
        )

    # =====================================================
    # SẮP XẾP
    # NGÀY HÓA ĐƠN MỚI NHẤT Ở TRÊN
    # =====================================================
    pos = (
        base_queryset
        .order_by(
            "-invoice__ngay_hd",
            "-po_number"
        )
    )

    # =====================================================
    # TỔNG TOÀN BỘ DỮ LIỆU SAU KHI LỌC
    #
    # Không phụ thuộc trang hiện tại
    # =====================================================
    total_import = 0
    total_export = 0

    for po in pos:

        invoice = po.invoice

        amount = invoice.tong_tien or 0

        if po.phan_loai_phieu in ["HH", "PN"]:
            total_import += amount

        elif po.phan_loai_phieu == "PX":
            total_export += amount

    # =====================================================
    # PHÂN TRANG
    # =====================================================

    per_page = request.GET.get(
        "per_page",
        "15"
    ).strip()

    page_number = request.GET.get(
        "page",
        "1"
    ).strip()

    # Luôn khởi tạo
    page_numbers = []

    if per_page.lower() == "all":

        paginator = None
        page_obj = None
        pos_page = pos

    else:

        try:
            per_page_int = int(per_page)

            if per_page_int <= 0:
                per_page_int = 15

        except (ValueError, TypeError):
            per_page_int = 15

        paginator = Paginator(
            pos,
            per_page_int
        )

        page_obj = paginator.get_page(
            page_number
        )

        pos_page = page_obj

        # =================================================
        # CÁC SỐ TRANG HIỂN THỊ
        # Ví dụ đang ở trang 5:
        # 3 4 5 6 7
        # =================================================

        current_page = page_obj.number
        total_pages = paginator.num_pages

        start_page = max(
            1,
            current_page - 2
        )

        end_page = min(
            total_pages,
            current_page + 2
        )

        page_numbers = range(
            start_page,
            end_page + 1
        )
    # =====================================================
    # TẠO DỮ LIỆU BÁO CÁO
    #
    # MỖI PHIẾU = 1 DÒNG
    # =====================================================
    report_data = []

    for po in pos_page:

        invoice = po.invoice

        amount = invoice.tong_tien or 0

        # =================================================
        # PHIẾU NHẬP
        # =================================================
        if po.phan_loai_phieu in ["HH", "PN"]:

            import_amount = amount
            export_amount = None

            # Nhà cung cấp
            partner_name = (
                getattr(
                    invoice,
                    "ten_nguoi_ban",
                    None
                )
                or getattr(
                    po,
                    "supplier",
                    None
                )
                or ""
            )

            partner_type = "Nhà cung cấp"

        # =================================================
        # PHIẾU XUẤT
        # =================================================
        elif po.phan_loai_phieu == "PX":

            import_amount = None
            export_amount = amount

            # Khách hàng
            partner_name = (
                getattr(
                    invoice,
                    "ten_nguoi_mua",
                    None
                )
                or ""
            )

            partner_type = "Khách hàng"

        else:
            continue

        # =================================================
        # THÊM DÒNG
        # =================================================
        report_data.append({
            "date": invoice.ngay_hd,

            "invoice_number": getattr(
                invoice,
                "so_hoa_don",
                ""
            ),

            "partner_name": partner_name,

            "partner_type": partner_type,

            "import_amount": import_amount,

            "export_amount": export_amount,

            "po_number": po.po_number,

            "activity_type": po.activity_type,
        })

    # DANH SÁCH NĂM TÀI CHÍNH
    #
    # Nguồn chính: Invoice.fiscal_year
    # Vì báo cáo đang lọc theo invoice__fiscal_year
    # =====================================================

    year_values = (
        Invoice.objects
        .exclude(fiscal_year__isnull=True)
        .values_list(
            "fiscal_year",
            flat=True
        )
        .distinct()
        .order_by("-fiscal_year")
    )

    year_list = [
        year
        for year in year_values
        if year
    ]

    # Đảm bảo năm hiện tại vẫn xuất hiện
    if (
        current_year != "all"
        and current_year not in year_list
    ):
        year_list.append(current_year)

    year_list = sorted(
        set(year_list),
        reverse=True
    )

    # =====================================================
    # QUERY PARAMS
    #
    # Dùng để giữ bộ lọc khi chuyển trang
    # =====================================================
    query_params = request.GET.copy()

    query_params.pop(
        "page",
        None
    )
    query_params = query_params.urlencode()
    # =====================================================
    # CONTEXT
    # =====================================================
    context = {
        "activity_types": activity_types,

        "selected_activity_type": selected_activity_type,

        "selected_year": selected_year,

        "start_date": start_date,
        "end_date": end_date,

        "report_data": report_data,

        "total_import": total_import,
        "total_export": total_export,

        "current_year": current_year,
        "year_list": year_list,

        "paginator": paginator,
        "page_obj": page_obj,
        "per_page": per_page,
        "page_numbers": page_numbers,

        "query_params": query_params,
    }   

    return render(
        request,
        "business_type_report.html",
        context
    )



from datetime import datetime

from django.http import HttpResponse
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

from invoice_reader_app.model_invoice import Invoice
from invoice_reader_app.models_activitytype import ActivityType
from invoice_reader_app.models_purchaseorder import PurchaseOrder


def export_business_type_report_excel(request):

    # =====================================================
    # NĂM TÀI CHÍNH
    # =====================================================

    current_year = request.session.get(
        "fiscal_year",
        datetime.now().year
    )

    year_param = request.GET.get(
        "year",
        ""
    ).strip()

    if year_param.lower() == "all":

        selected_year = "all"

    elif year_param:

        try:
            selected_year = int(year_param)
            current_year = selected_year

        except (ValueError, TypeError):

            selected_year = current_year

    else:

        selected_year = current_year

    # =====================================================
    # LOẠI HÌNH KINH DOANH
    # =====================================================

    selected_activity_type = request.GET.get(
        "activity_type",
        ""
    ).strip()

    # =====================================================
    # NGÀY
    # =====================================================

    start_date = request.GET.get(
        "start_date",
        ""
    ).strip()

    end_date = request.GET.get(
        "end_date",
        ""
    ).strip()

    # =====================================================
    # QUERY
    # =====================================================

    base_queryset = (
        PurchaseOrder.objects
        .select_related(
            "invoice",
            "activity_type"
        )
        .filter(
            invoice__ngay_hd__isnull=False,
            phan_loai_phieu__in=[
                "HH",
                "PN",
                "PX",
            ],
        )
    )

    # =====================================================
    # LỌC NĂM
    # =====================================================

    if selected_year != "all":

        base_queryset = base_queryset.filter(
            invoice__fiscal_year=selected_year
        )

    # =====================================================
    # LỌC LOẠI HÌNH
    # =====================================================

    if selected_activity_type:

        base_queryset = base_queryset.filter(
            activity_type_id=selected_activity_type
        )

    # =====================================================
    # LỌC NGÀY HÓA ĐƠN
    # =====================================================

    if start_date:

        base_queryset = base_queryset.filter(
            invoice__ngay_hd__gte=start_date
        )

    if end_date:

        base_queryset = base_queryset.filter(
            invoice__ngay_hd__lte=end_date
        )

    # =====================================================
    # SẮP XẾP
    # NGÀY MỚI NHẤT Ở TRÊN
    # =====================================================

    pos = (
        base_queryset
        .order_by(
            "-invoice__ngay_hd",
            "-po_number"
        )
    )

    # =====================================================
    # TẠO EXCEL
    # =====================================================

    wb = Workbook()
    ws = wb.active
    ws.title = "Báo cáo kinh doanh"

    # =====================================================
    # TIÊU ĐỀ
    # =====================================================

    ws.merge_cells(
        "A1:G1"
    )

    ws["A1"] = "BÁO CÁO KINH DOANH"

    ws["A1"].font = Font(
        bold=True,
        size=16
    )

    ws["A1"].alignment = Alignment(
        horizontal="center",
        vertical="center"
    )

    ws.row_dimensions[1].height = 28

    # =====================================================
    # THÔNG TIN BỘ LỌC
    # =====================================================

    if selected_year == "all":
        year_text = "All"
    else:
        year_text = str(selected_year)

    activity_text = "Tất cả"

    if selected_activity_type:

        activity = (
            ActivityType.objects
            .filter(
                id=selected_activity_type
            )
            .first()
        )

        if activity:
            activity_text = (
                f"{activity.code} - {activity.name}"
            )

    ws["A2"] = "Năm tài chính"
    ws["B2"] = year_text

    ws["A3"] = "Loại hình kinh doanh"
    ws["B3"] = activity_text

    ws["A4"] = "Từ ngày"
    ws["B4"] = start_date or "Tất cả"

    ws["D4"] = "Đến ngày"
    ws["E4"] = end_date or "Tất cả"

    for cell in [
        "A2",
        "A3",
        "A4",
        "D4",
    ]:
        ws[cell].font = Font(
            bold=True
        )

    # =====================================================
    # HEADER
    # =====================================================

    header_row = 6

    headers = [
        "TT",
        "Ngày hóa đơn",
        "Số hóa đơn",
        "Nhà cung cấp / Khách hàng",
        "Loại hình",
        "Nhập",
        "Xuất",
    ]

    header_fill = PatternFill(
        "solid",
        fgColor="4F81BD"
    )

    header_font = Font(
        bold=True,
        color="FFFFFF"
    )

    thin = Side(
        style="thin",
        color="BFBFBF"
    )

    border = Border(
        left=thin,
        right=thin,
        top=thin,
        bottom=thin
    )

    for col, title in enumerate(
        headers,
        start=1
    ):

        cell = ws.cell(
            row=header_row,
            column=col,
            value=title
        )

        cell.fill = header_fill
        cell.font = header_font

        cell.alignment = Alignment(
            horizontal="center",
            vertical="center",
            wrap_text=True
        )

        cell.border = border

    # =====================================================
    # DỮ LIỆU
    # =====================================================

    row_number = header_row + 1
    stt = 1

    total_import = 0
    total_export = 0

    for po in pos:

        invoice = po.invoice

        amount = invoice.tong_tien or 0

        # ---------------------------------------------
        # NHẬP
        # ---------------------------------------------

        if po.phan_loai_phieu in [
            "HH",
            "PN",
        ]:

            import_amount = amount
            export_amount = None

            supplier = getattr(
                po,
                "supplier",
                None
            )

            partner_name = (
                getattr(
                    invoice,
                    "ten_nguoi_ban",
                    None
                )
                or (
                    str(supplier)
                    if supplier
                    else ""
                )
                or ""
            )

            partner_type = "Nhà cung cấp"

            total_import += amount

        # ---------------------------------------------
        # XUẤT
        # ---------------------------------------------

        elif po.phan_loai_phieu == "PX":

            import_amount = None
            export_amount = amount

            partner_name = (
                getattr(
                    invoice,
                    "ten_nguoi_mua",
                    None
                )
                or ""
            )

            partner_type = "Khách hàng"

            total_export += amount

        else:
            continue

        # ---------------------------------------------
        # LOẠI HÌNH
        # ---------------------------------------------

        if po.activity_type:

            activity_name = (
                f"{po.activity_type.code} - "
                f"{po.activity_type.name}"
            )

        else:

            activity_name = ""

        # ---------------------------------------------
        # GHI DÒNG
        # ---------------------------------------------

        values = [
            stt,
            invoice.ngay_hd,
            getattr(invoice, "so_hoa_don", "") or "",
            str(partner_name) if partner_name else "",
            str(activity_name) if activity_name else "",
            import_amount,
            export_amount,
        ]

        for col, value in enumerate(
            values,
            start=1
        ):

            cell = ws.cell(
                row=row_number,
                column=col,
                value=value
            )

            cell.border = border

            if col in [
                1,
                2,
                3,
            ]:
                cell.alignment = Alignment(
                    horizontal="center",
                    vertical="center"
                )

            elif col in [
                6,
                7,
            ]:
                cell.alignment = Alignment(
                    horizontal="right"
                )

            else:
                cell.alignment = Alignment(
                    vertical="center"
                )

        # Ngày
        ws.cell(
            row=row_number,
            column=2
        ).number_format = "dd/mm/yyyy"

        # Tiền
        ws.cell(
            row=row_number,
            column=6
        ).number_format = '#,##0'

        ws.cell(
            row=row_number,
            column=7
        ).number_format = '#,##0'

        row_number += 1
        stt += 1

    # =====================================================
    # TỔNG CỘNG
    # =====================================================

    total_row = row_number

    ws.cell(
        row=total_row,
        column=1,
        value="TỔNG CỘNG"
    )

    ws.merge_cells(
        start_row=total_row,
        start_column=1,
        end_row=total_row,
        end_column=5
    )

    ws.cell(
        row=total_row,
        column=6,
        value=total_import
    )

    ws.cell(
        row=total_row,
        column=7,
        value=total_export
    )

    for col in range(1, 8):

        cell = ws.cell(
            row=total_row,
            column=col
        )

        cell.font = Font(
            bold=True
        )

        cell.fill = PatternFill(
            "solid",
            fgColor="D9EAD3"
        )

        cell.border = border

    ws.cell(
        row=total_row,
        column=1
    ).alignment = Alignment(
        horizontal="right"
    )

    ws.cell(
        row=total_row,
        column=6
    ).number_format = '#,##0'

    ws.cell(
        row=total_row,
        column=7
    ).number_format = '#,##0'

    # =====================================================
    # ĐỘ RỘNG CỘT
    # =====================================================

    column_widths = {
        "A": 8,
        "B": 15,
        "C": 15,
        "D": 40,
        "E": 30,
        "F": 20,
        "G": 20,
    }

    for column, width in column_widths.items():

        ws.column_dimensions[
            column
        ].width = width

    # =====================================================
    # FREEZE HEADER
    # =====================================================

    ws.freeze_panes = "A7"

    # =====================================================
    # FILTER
    # =====================================================

    if row_number > header_row + 1:

        ws.auto_filter.ref = (
            f"A{header_row}:G{row_number - 1}"
        )

    # =====================================================
    # RESPONSE
    # =====================================================

    response = HttpResponse(
        content_type=(
            "application/vnd.openxmlformats-"
            "officedocument.spreadsheetml.sheet"
        )
    )

    response[
        "Content-Disposition"
    ] = (
        'attachment; '
        'filename="bao_cao_kinh_doanh.xlsx"'
    )

    wb.save(response)

    return response